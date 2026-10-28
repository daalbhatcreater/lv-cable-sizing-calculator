"""Size every circuit in a CSV and write a cable schedule.

    python -m cablesizer examples/feeders.csv -o results/cable_schedule.xlsx
"""

import argparse
import csv

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from .sizing import Circuit, size_circuit


def read_circuits(path):
    circuits = []
    with open(path, newline="") as fh:
        for row in csv.DictReader(fh):
            kwargs = {}
            for key, value in row.items():
                value = value.strip()
                if value == "":
                    continue
                if key in ("name", "insulation", "method", "arrangement"):
                    kwargs[key] = value
                elif key in ("phases", "grouped", "runs"):
                    kwargs[key] = int(value)
                else:
                    kwargs[key] = float(value)
            circuits.append(Circuit(**kwargs))
    return circuits


COLUMNS = [
    ("Circuit", 22), ("Ib (A)", 9), ("In (A)", 9), ("Cable", 14),
    ("It (A)", 9), ("Ca", 7), ("Cg", 7), ("Iz (A)", 9), ("VD (%)", 9),
    ("Start VD (%)", 12), ("S min SC (mm2)", 14), ("Governed by", 24), ("Notes", 30),
]


def write_schedule(circuits, results, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cable schedule"
    head_fill = PatternFill("solid", fgColor="1F3864")
    for col, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = width
    for r, (c, res) in enumerate(zip(circuits, results), start=2):
        values = [
            res.name, round(res.ib_a, 1), res.device_a,
            f"{res.designation} {c.insulation}", res.it_a, res.ca, res.cg,
            round(res.iz_a, 1), round(res.vd_pct, 2),
            round(res.start_vd_pct, 2) if res.start_vd_pct is not None else "",
            round(res.s_min_sc_mm2, 1) if res.s_min_sc_mm2 else "",
            res.governed_by, "; ".join(res.notes),
        ]
        for col, v in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=v)
    ws.freeze_panes = "B2"
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("csv")
    ap.add_argument("-o", "--output", default="cable_schedule.xlsx")
    args = ap.parse_args()

    circuits = read_circuits(args.csv)
    results = [size_circuit(c) for c in circuits]
    print(f"{'Circuit':22} {'Ib':>7} {'In':>6} {'Cable':>14} {'Iz':>7} {'VD%':>6}  governed by")
    for r in results:
        print(f"{r.name:22} {r.ib_a:7.1f} {r.device_a:6g} {r.designation:>14} "
              f"{r.iz_a:7.1f} {r.vd_pct:6.2f}  {r.governed_by}")
    write_schedule(circuits, results, args.output)
    print(f"\nwritten {args.output}")


if __name__ == "__main__":
    main()
